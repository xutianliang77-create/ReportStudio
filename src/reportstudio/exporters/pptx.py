from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True, slots=True)
class PptxArtifact:
    path: str


def export_pptx(out_dir: str, file_name: str, title: str, kpis: dict[str, float]) -> PptxArtifact:
    from pptx import Presentation

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / file_name

    prs = Presentation()

    # cover
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = title
    s1.placeholders[1].text = "ReportStudio Community"

    # KPI slide
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "KPI Summary"
    tf = s2.placeholders[1].text_frame
    tf.clear()
    for i, (k, v) in enumerate(list(kpis.items())[:12]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{k}: {v:,.2f}" if isinstance(v, float) else f"{k}: {v}"
        p.level = 0

    prs.save(str(path))
    return PptxArtifact(path=str(path))
