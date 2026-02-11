from __future__ import annotations

from pathlib import Path

from reportstudio.exporters.pdf import export_pdf
from reportstudio.exporters.pptx import export_pptx


def test_export_pdf_structured_non_empty(tmp_path: Path) -> None:
    out = export_pdf(
        str(tmp_path),
        "brief.pdf",
        title="Test Report",
        spec={
            "time_range": "2026-01-01..2026-02-01",
            "topn": 5,
            "date_column": "date",
            "number_columns": ["sales"],
            "dimension_columns": ["channel"],
        },
        kpis={"sales": 1234.5, "orders": 42, "sales_delta": 0.12},
        trend_rows=[
            {"period": "2026-01", "sales": 100.0},
            {"period": "2026-02", "sales": 200.0},
        ],
        breakdowns=[
            {
                "dim": "channel",
                "measure": "sales",
                "rows": [
                    {"channel": "A", "sales": 10.0},
                    {"channel": "B", "sales": 9.0},
                ],
            }
        ],
        highlights=["H1", "H2"],
        risks=["R1"],
        actions=["A1"],
        warnings=["W1"],
    )

    p = Path(out.path)
    assert p.exists()
    assert p.stat().st_size > 500

    # We disable compression in the exporter, so text markers are visible in bytes.
    b = p.read_bytes()
    for marker in [
        b"Community v1 Report",
        b"Spec",
        b"KPI Summary",
        b"Trend Summary",
        b"Breakdown Summary",
        b"Warnings & Limitations",
    ]:
        assert marker in b

    # 6 pages => 6 occurrences of /Type /Page (best-effort check)
    assert b.count(b"/Type /Page") >= 6


def test_export_pptx_structured_slide_count(tmp_path: Path) -> None:
    out = export_pptx(
        str(tmp_path),
        "deck.pptx",
        title="Test Report",
        spec={"time_range": "2026-01-01..2026-02-01", "topn": 5},
        kpis={"sales": 1234.5, "orders": 42},
        trend_rows=[
            {"period": "2026-01", "sales": 100.0},
            {"period": "2026-02", "sales": 200.0},
        ],
        breakdowns=[
            {
                "dim": "channel",
                "measure": "sales",
                "rows": [
                    {"channel": "A", "sales": 10.0},
                    {"channel": "B", "sales": 9.0},
                ],
            }
        ],
        highlights=["H1"],
        risks=["R1"],
        actions=["A1"],
    )

    p = Path(out.path)
    assert p.exists()
    assert p.stat().st_size > 5000

    from pptx import Presentation

    prs = Presentation(str(p))
    assert len(prs.slides) == 5

    titles = [s.shapes.title.text for s in prs.slides]
    assert titles[1:] == [
        "KPI Summary",
        "Trend Summary",
        "Breakdown Summary",
        "Risks & Actions",
    ]
