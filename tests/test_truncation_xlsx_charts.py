from __future__ import annotations

import json
import subprocess
import sys
from datetime import date, timedelta
from pathlib import Path

import pandas as pd


def test_large_trend_triggers_truncation_warning_and_xlsx_note_and_chart(tmp_path: Path) -> None:
    # Create >5000 daily periods so XLSX exporter must truncate.
    start = date(2000, 1, 1)
    days = 6001
    dts = [(start + timedelta(days=i)).isoformat() for i in range(days)]

    df = pd.DataFrame(
        {
            "dt": dts,
            "revenue": list(range(days)),
            # include a dim so pipeline can build breakdowns, but keep it tiny/constant
            "region": ["A"] * days,
        }
    )

    inp = tmp_path / "big.csv"
    df.to_csv(inp, index=False)

    out_dir = tmp_path / "out"

    cmd = [
        sys.executable,
        "-m",
        "reportstudio.cli.main",
        "--file",
        str(inp),
        "--prompt",
        "big trend",
        "--out-dir",
        str(out_dir),
        "--grain",
        "day",
        "--time-range",
        "2000-01-01..2020-01-01",
    ]
    p = subprocess.run(cmd, check=True, capture_output=True, text=True)
    payload = json.loads(p.stdout)

    # CLI-level warnings should include truncation marker
    warnings = payload.get("warnings", [])
    assert any("Trend table" in w and "truncated" in w for w in warnings)

    # XLSX should include note in Trend sheet and have a chart object.
    xlsx_paths = [a["path"] for a in payload["artifacts"] if a["format"] == "xlsx"]
    assert len(xlsx_paths) == 1
    xlsx_path = Path(xlsx_paths[0])
    assert xlsx_path.exists()

    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    assert {"Summary", "Trend"}.issubset(set(wb.sheetnames))

    ts = wb["Trend"]

    # Note row present
    assert ts["A1"].value == "NOTE"
    assert isinstance(ts["B1"].value, str)
    assert "truncated" in ts["B1"].value.lower()

    # Header row should include period
    assert ts["A3"].value in {"period", "dt"}

    # Data rows truncated to 5000
    # Row1 note, Row2 blank, Row3 header, then 5000 data rows => 5003 total
    assert ts.max_row == 5003

    # Chart objects exist (best-effort; we add a line chart)
    assert len(getattr(ts, "_charts", [])) > 0
