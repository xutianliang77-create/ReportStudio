from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from reportstudio.exporters.xlsx import export_xlsx


def test_export_xlsx_has_professional_layout_and_charts(tmp_path: Path) -> None:
    kpis = {
        "orders": 1200.0,
        "gmv": 98765.43,
        "conversion_rate": 0.0345,
    }
    warnings = ["Schema guessed; results may vary"]
    spec = {
        "title": "Community v1.0",
        "time_range": "2026-01-01..2026-02-01",
        "topn": 5,
        "xlsx_limits": {"trend_max_rows": 5, "breakdown_max_rows": 5, "contributors_max_rows": 5},
    }

    trend_rows = [
        {"period": "2026-01", "orders": 100, "conversion_rate": 0.03},
        {"period": "2026-02", "orders": 120, "conversion_rate": 0.035},
        {"period": "2026-03", "orders": 110, "conversion_rate": 0.033},
    ]

    breakdowns = [
        {
            "dim": "community",
            "measure": "orders",
            "rows": [
                {"community": "A", "orders": 50, "share": 0.25},
                {"community": "B", "orders": 100, "share": 0.50},
                {"community": "C", "orders": 50, "share": 0.25},
            ],
        }
    ]

    art = export_xlsx(
        str(tmp_path),
        "out.xlsx",
        kpis=kpis,
        warnings=warnings,
        spec=spec,
        trend_rows=trend_rows,
        breakdowns=breakdowns,
    )

    wb = load_workbook(art.path)
    assert set(wb.sheetnames) >= {"Summary", "Trend", "Breakdowns"}

    ws = wb["Summary"]
    assert ws["A1"].value == "Community v1.0"
    assert ws.freeze_panes == "A6"

    ts = wb["Trend"]
    assert ts.freeze_panes == "A2"
    assert len(ts._charts) >= 1

    bs = wb["Breakdowns"]
    assert bs.freeze_panes == "A2"
    # bar chart should be placed on the breakdowns sheet
    assert len(bs._charts) >= 1


def test_export_xlsx_truncates_large_tables_and_notes_it(tmp_path: Path) -> None:
    kpis = {"orders": 1.0}
    warnings: list[str] = []
    spec = {
        "time_range": "(none)",
        "topn": 5,
        "xlsx_limits": {"trend_max_rows": 3, "breakdown_max_rows": 2, "contributors_max_rows": 2},
    }

    trend_rows = [{"period": f"p{i}", "orders": i} for i in range(10)]
    breakdowns = [
        {
            "dim": "d",
            "measure": "m",
            "rows": [{"d": f"x{i}", "m": i} for i in range(10)],
        }
    ]

    art = export_xlsx(
        str(tmp_path),
        "out.xlsx",
        kpis=kpis,
        warnings=warnings,
        spec=spec,
        trend_rows=trend_rows,
        breakdowns=breakdowns,
    )

    wb = load_workbook(art.path)
    ws = wb["Summary"]

    # Notes are written back as an additional block at the bottom of Summary.
    values = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert any("Trend table" in v and "truncated" in v for v in values)
    assert any("Breakdown(d)" in v and "truncated" in v for v in values)
